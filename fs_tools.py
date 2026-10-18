# Assignment 1, Part A: four file tools
#
#   read_file(filepath)                 read a PDF, DOCX or TXT file
#   list_files(directory, extension)    list the files in a folder
#   write_file(filepath, content)       save text to a file
#   search_in_file(filepath, keyword)   find a word in a file, with the text around it
#
# Every tool returns a plain dictionary (or a list of dictionaries).
# When something goes wrong, the tool does NOT crash. It returns
#   {"success": False, "error": "what went wrong"}
# so the program, and later the LLM, can carry on and explain the problem.
#
# Try it with:  python3 fs_tools.py

import os
from datetime import datetime

from docx import Document          # reads Word .docx files
from pypdf import PdfReader        # reads PDF files

SUPPORTED_TYPES = [".pdf", ".docx", ".txt"]


def get_file_info(filepath):
    """Return the name, size and last-modified date of a file."""
    size_in_bytes = os.path.getsize(filepath)

    # getmtime gives the time as a number of seconds since 1 January 1970.
    # datetime turns that number into a readable date like 2026-10-16T18:26:52
    modified_seconds = os.path.getmtime(filepath)
    modified_date = datetime.fromtimestamp(modified_seconds).isoformat(timespec="seconds")

    return {
        "name": os.path.basename(filepath),               # "resumes/a.pdf" -> "a.pdf"
        "path": filepath,
        "extension": os.path.splitext(filepath)[1].lower(),   # "a.PDF" -> ".pdf"
        "size_bytes": size_in_bytes,
        "modified": modified_date,
    }


# ---------------------------------------------------------------------------
# Tool 1: read_file
# ---------------------------------------------------------------------------
def read_file(filepath):
    """Read a PDF, DOCX or TXT file. Return its text and some facts about it."""

    # Check the problems we can see before opening the file
    if not os.path.exists(filepath):
        return {"success": False, "error": "File not found: " + filepath}

    if not os.path.isfile(filepath):
        return {"success": False, "error": "Not a file (it is a folder): " + filepath}

    extension = os.path.splitext(filepath)[1].lower()
    if extension not in SUPPORTED_TYPES:
        return {"success": False, "error": "Unsupported file type " + extension + ". Use PDF, DOCX or TXT."}

    # Each file type needs its own reader. A damaged or password-protected file
    # makes the reader fail, so we catch that and return an error instead.
    try:
        if extension == ".pdf":
            reader = PdfReader(filepath)
            text = ""
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:                 # a scanned page gives no text, so skip it
                    text = text + page_text + "\n"

        elif extension == ".docx":
            document = Document(filepath)
            text = ""
            for paragraph in document.paragraphs:
                text = text + paragraph.text + "\n"
            # Resumes often put skills in a table, which paragraphs do not include
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text = text + cell.text + " | "
                    text = text + "\n"

        else:  # .txt
            with open(filepath, "r", encoding="utf-8", errors="replace") as file:
                text = file.read()

    except Exception as error:
        return {"success": False, "error": "Could not read " + filepath + ": " + str(error)}

    text = text.strip()                   # remove empty lines at the start and end

    metadata = get_file_info(filepath)
    metadata["characters"] = len(text)
    metadata["words"] = len(text.split())  # split() cuts the text at every space or new line

    if text == "":
        metadata["warning"] = "No text found. The file may be a scanned image."

    return {"success": True, "content": text, "metadata": metadata}


# ---------------------------------------------------------------------------
# Tool 2: list_files
# ---------------------------------------------------------------------------
def list_files(directory, extension=None):
    """List the files in a folder. Optionally keep only one type, such as ".pdf"."""

    if not os.path.isdir(directory):
        # The brief says this tool returns a list, so the error goes inside a list too
        return [{"error": "Folder not found: " + directory}]

    # Accept "pdf", ".pdf", "PDF" or ".PDF" and turn them all into ".pdf"
    if extension:
        extension = extension.lower()
        if not extension.startswith("."):
            extension = "." + extension

    files = []
    for name in sorted(os.listdir(directory)):      # sorted = alphabetical order
        full_path = os.path.join(directory, name)

        if not os.path.isfile(full_path):            # skip sub-folders
            continue

        file_extension = os.path.splitext(name)[1].lower()
        if extension and file_extension != extension:
            continue                                  # skip files of other types

        files.append(get_file_info(full_path))

    return files


# ---------------------------------------------------------------------------
# Tool 3: write_file
# ---------------------------------------------------------------------------
def write_file(filepath, content):
    """Save text to a file. Creates missing folders. Replaces the file if it exists."""
    try:
        folder = os.path.dirname(filepath)         # "summaries/a.txt" -> "summaries"
        if folder:
            os.makedirs(folder, exist_ok=True)     # create it if it does not exist yet

        already_existed = os.path.exists(filepath)

        with open(filepath, "w", encoding="utf-8") as file:
            file.write(content)

    except Exception as error:
        return {"success": False, "error": "Could not write " + filepath + ": " + str(error)}

    return {
        "success": True,
        "filepath": filepath,
        # Characters and bytes differ: "₹" is 1 character but 3 bytes in UTF-8
        "bytes_written": len(content.encode("utf-8")),
        "overwritten": already_existed,
    }


# ---------------------------------------------------------------------------
# Tool 4: search_in_file
# ---------------------------------------------------------------------------
def search_in_file(filepath, keyword, context_chars=60):
    """Find every place a keyword appears (any upper/lower case), with the text around it."""

    if keyword is None or keyword.strip() == "":
        return {"success": False, "error": "The keyword must not be empty."}

    # Reuse read_file so we do not repeat the PDF/DOCX/TXT reading code
    result = read_file(filepath)
    if not result["success"]:
        return result

    text = result["content"]
    keyword = keyword.strip()

    # Compare lowercase copies, so "Python", "python" and "PYTHON" all match
    lower_text = text.lower()
    lower_keyword = keyword.lower()

    matches = []
    position = lower_text.find(lower_keyword)       # -1 means "not found"

    while position != -1:
        # Take up to 60 characters before and after the match.
        # max(0, ...) stops us going before the start of the text,
        # min(len(text), ...) stops us going past the end.
        start = max(0, position - context_chars)
        end = min(len(text), position + len(keyword) + context_chars)

        snippet = text[start:end]
        snippet = " ".join(snippet.split())          # turn new lines and double spaces into single spaces

        matches.append({
            "position": position,
            "match": text[position:position + len(keyword)],   # the word as written in the file
            "context": snippet,
        })

        # Look for the next match, starting just after this one
        position = lower_text.find(lower_keyword, position + 1)

    return {
        "success": True,
        "filepath": filepath,
        "keyword": keyword,
        "match_count": len(matches),
        "matches": matches,
    }


# This part only runs when you start this file directly: python3 fs_tools.py
if __name__ == "__main__":
    print("list_files:")
    for file_info in list_files("resumes"):
        print(" ", file_info)

    print("\nread_file metadata:")
    print(" ", read_file("resumes/resume_john_doe.pdf")["metadata"])

    print("\nsearch_in_file:")
    search = search_in_file("resumes/resume_sara_khan.docx", "python")
    print("  matches found:", search["match_count"])
    for match in search["matches"]:
        print("  ...", match["context"], "...")
